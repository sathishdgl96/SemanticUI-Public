"""Export real data from the real account, and open the result."""

import io
import os
from datetime import datetime, timezone

import openpyxl
import pytest

pytestmark = [
    pytest.mark.integration,
    pytest.mark.skipif(
        not os.environ.get("SEMANTICUI_IT_ACCOUNT"),
        reason="SEMANTICUI_IT_* env vars not set",
    ),
]


def _connect():
    from app.snowflake import connect as sf_connect

    return sf_connect.connect_dev(
        account=os.environ["SEMANTICUI_IT_ACCOUNT"],
        user=os.environ["SEMANTICUI_IT_USER"],
        authenticator="password",
        password=os.environ["SEMANTICUI_IT_PASSWORD"],
    )


def _view():
    return (
        os.environ["SEMANTICUI_IT_DATABASE"],
        os.environ["SEMANTICUI_IT_SCHEMA"],
        os.environ["SEMANTICUI_IT_VIEW"],
    )


def test_a_real_export_opens_as_a_workbook():
    """Bytes openpyxl refuses are bytes Excel would refuse too."""
    from app.export.workbook import SheetData, build_workbook
    from app.semantic.discovery import describe_semantic_view

    conn = _connect()
    try:
        db, schema, view = _view()
        detail = describe_semantic_view(conn, db, schema, view)
        dim = detail["dimensions"][0]

        from app.semantic.query import SemanticQueryRequest, build_semantic_sql
        from app.snowflake.gateway import run_query

        request = SemanticQueryRequest.model_validate(
            {
                "database": db,
                "schema": schema,
                "view": view,
                "dimensions": [f"{dim['table']}.{dim['name']}"],
                "limit": 50,
            }
        )
        sql, params, limit = build_semantic_sql(detail, request, max_rows=10000)
        result = run_query(conn, sql, max_rows=limit, params=params)

        data = build_workbook(
            "Integration export",
            f"{db}.{schema}.{view}",
            os.environ["SEMANTICUI_IT_USER"],
            [SheetData(title=dim["name"], columns=result.columns, rows=result.rows)],
            generated_at=datetime.now(timezone.utc),
        )
        workbook = openpyxl.load_workbook(io.BytesIO(data))
        assert workbook.sheetnames[0] == "Summary"
        assert len(workbook.sheetnames) == 2
        data_sheet = workbook[workbook.sheetnames[1]]
        assert data_sheet.cell(row=1, column=1).value == dim["name"]
        assert data_sheet.max_row > 1, "the real view returned no rows"
        print(
            f"\n[export] {len(data)} bytes, sheets={workbook.sheetnames}, "
            f"rows={data_sheet.max_row - 1}"
        )
    finally:
        conn.close()


def test_the_copyable_sql_runs_against_the_real_account():
    """The statement handed to Excel has to be one Snowflake actually accepts.

    Run here rather than trusted: an inlined literal that Snowflake rejects
    would be discovered by the user, in Excel, with no error message from us.
    """
    from app.export.literals import build_literal_sql
    from app.semantic.discovery import describe_semantic_view
    from app.semantic.query import SemanticQueryRequest

    conn = _connect()
    try:
        db, schema, view = _view()
        detail = describe_semantic_view(conn, db, schema, view)
        dim = detail["dimensions"][0]
        ref = f"{dim['table']}.{dim['name']}"

        request = SemanticQueryRequest.model_validate(
            {
                "database": db,
                "schema": schema,
                "view": view,
                "dimensions": [ref],
                "limit": 5,
                "filters": [
                    {"id": "f", "field": ref, "op": "isNot", "values": ["O'Brien"]}
                ],
            }
        )
        sql = build_literal_sql(detail, request)
        assert "?" not in sql
        assert "'O''Brien'" in sql

        cur = conn.cursor()
        try:
            cur.execute(sql)
            rows = cur.fetchmany(5)
            print(f"\n[connect-sql] accepted by Snowflake, {len(rows)} row(s)")
        finally:
            cur.close()
    finally:
        conn.close()

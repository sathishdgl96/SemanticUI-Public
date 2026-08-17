"""A workbook that carries its own live connection.

This edits a zip of XML parts by hand, which is exactly the kind of code that
produces a file Excel offers to "repair". Excel is not available here, so the
tests below check every structural property that CAN be checked without it:
the archive opens, every part is well-formed XML, every part has a content
type, every relationship resolves to a part that exists, and openpyxl can read
the result. That is not proof Excel accepts it, and the manual-pass note says
so -- but it is what stands between a plausible file and a corrupt one.
"""

import io
import re
import zipfile
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook

from app.export.live import (
    DRIVER,
    LiveSheet,
    _column_letter,
    _table_name,
    _unique_columns,
    add_live_connections,
    connection_string,
)
from app.export.workbook import SheetData, build_workbook

SQL = 'SELECT * FROM SEMANTIC_VIEW("D"."S"."V" DIMENSIONS "C"."REGION")'


def base_workbook(rows=3):
    return build_workbook(
        "Sales",
        "D.S.V",
        "ALICE",
        [
            SheetData(
                title="Revenue by region",
                columns=[{"name": "REGION"}, {"name": "TOTAL"}],
                rows=[[f"R{i}", i] for i in range(rows)],
            )
        ],
        generated_at=datetime(2026, 8, 17, tzinfo=timezone.utc),
    )


def live(xlsx=None, sheets=None, **kwargs):
    options = {
        "account": "XRIIEIM-EH01350",
        "database": "SEMANTIC_DEMO",
        "schema": "TPCH",
    }
    options.update(kwargs)
    return add_live_connections(
        xlsx if xlsx is not None else base_workbook(),
        sheets
        if sheets is not None
        else [LiveSheet(title="Revenue by region", sql=SQL, columns=["REGION", "TOTAL"], rows=3)],
        **options,
    )


def parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        return {i.filename: z.read(i.filename) for i in z.infolist()}


class TestStructure:
    """The properties that separate a workbook from a repair prompt."""

    def test_every_part_is_well_formed_xml(self):
        from xml.etree import ElementTree

        for name, blob in parts(live()).items():
            if name.endswith(".xml") or name.endswith(".rels"):
                ElementTree.fromstring(blob)  # raises on malformed

    def test_every_part_has_a_content_type(self):
        archive = parts(live())
        types = archive["[Content_Types].xml"].decode()
        defaults = set(re.findall(r'Default Extension="([^"]+)"', types))
        overrides = set(re.findall(r'Override PartName="([^"]+)"', types))
        for name in archive:
            if name == "[Content_Types].xml":
                continue
            extension = name.rsplit(".", 1)[-1]
            assert f"/{name}" in overrides or extension in defaults, name

    def test_every_relationship_points_at_a_part_that_exists(self):
        # A dangling target is the single most common way a hand-built xlsx
        # opens as "unreadable content".
        archive = parts(live())
        for name, blob in archive.items():
            if not name.endswith(".rels"):
                continue
            # A rels part sits in `<folder>/_rels/`, and the package root's
            # own `_rels/.rels` has no folder at all.
            folder = name.rsplit("/_rels/", 1)[0] if "/_rels/" in name else ""
            for target, mode in re.findall(
                r'Target="([^"]+)"(?:\s+TargetMode="([^"]+)")?', blob.decode()
            ):
                if mode == "External" or target.startswith("http"):
                    continue
                resolved = _resolve(folder, target)
                assert resolved in archive, f"{name} -> {target}"

    def test_openpyxl_can_still_read_it(self):
        book = load_workbook(io.BytesIO(live()))
        assert "Summary" in book.sheetnames
        sheet = book["Revenue by region"]
        assert sheet["A1"].value == "REGION"

    def test_table_parts_is_the_last_child_of_the_worksheet(self):
        # The schema is a sequence; anywhere else and Excel rejects the part.
        archive = parts(live())
        sheet = next(
            blob.decode()
            for name, blob in archive.items()
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml")
            and "tableParts" in blob.decode()
        )
        assert sheet.rstrip().endswith("</tableParts></worksheet>")


class TestConnection:
    def test_it_carries_the_server_the_driver_and_the_query(self):
        connections = parts(live())["xl/connections.xml"].decode()
        assert DRIVER in connections
        assert "XRIIEIM-EH01350.snowflakecomputing.com" in connections
        assert "SEMANTIC_VIEW" in connections
        assert 'commandType="2"' in connections

    def test_the_query_is_attribute_escaped_rather_than_breaking_the_part(self):
        sheets = [
            LiveSheet(
                title="Revenue by region",
                sql="SELECT * WHERE name = 'a\"b' AND x < 5",
                columns=["REGION", "TOTAL"],
                rows=3,
            )
        ]
        connections = parts(live(sheets=sheets))["xl/connections.xml"].decode()
        assert "&quot;" in connections or "&#34;" in connections
        assert "&lt;" in connections

    def test_no_credential_is_written(self):
        # The product's premise is that every query runs on the caller's own
        # credentials; a workbook carrying one hands it to everybody it is
        # forwarded to. There is nowhere here to put a secret, and this test
        # is what keeps it that way.
        blob = live().lower()
        for word in (b"password", b"pwd=", b"uid=", b"token=", b"authenticator"):
            assert word not in blob

    def test_a_warehouse_appears_only_when_asked_for(self):
        assert "WAREHOUSE" not in connection_string("a", "b", "c", "")
        assert "WAREHOUSE=BI_WH" in connection_string("a", "b", "c", "BI_WH")


class TestWhatItSkips:
    def test_an_empty_sheet_gets_no_table(self):
        # A table needs a range. An empty one is the shape Excel repairs.
        sheets = [LiveSheet(title="Revenue by region", sql=SQL, columns=["A"], rows=0)]
        assert "xl/connections.xml" not in parts(live(sheets=sheets))

    def test_a_sheet_that_is_not_in_the_workbook_is_ignored(self):
        sheets = [LiveSheet(title="Nowhere", sql=SQL, columns=["A"], rows=2)]
        assert live(sheets=sheets) == base_workbook()

    def test_no_sheets_leaves_the_workbook_byte_for_byte(self):
        original = base_workbook()
        assert add_live_connections(
            original, [], account="a", database="b", schema="c"
        ) == original


class TestNaming:
    @pytest.mark.parametrize(
        "index,letter", [(0, "A"), (25, "Z"), (26, "AA"), (27, "AB"), (51, "AZ")]
    )
    def test_column_letters(self, index, letter):
        assert _column_letter(index) == letter

    def test_a_table_name_is_a_name_excel_accepts(self):
        assert _table_name("Revenue by region", set()) == "Revenue_by_region"
        assert _table_name("2026 sales", set()).startswith("_")

    def test_table_names_are_unique_within_a_workbook(self):
        taken: set[str] = set()
        assert _table_name("Revenue", taken) == "Revenue"
        assert _table_name("Revenue", taken) == "Revenue_2"

    def test_duplicate_columns_are_made_distinct(self):
        # Excel refuses a table with two columns of the same name, and two
        # entities can each have a NAME column.
        assert _unique_columns(["NAME", "NAME", ""]) == ["NAME", "NAME_2", "Column3"]


def _resolve(folder: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    segments = folder.split("/") if folder else []
    for piece in target.split("/"):
        if piece == "..":
            segments.pop()
        elif piece not in ("", "."):
            segments.append(piece)
    return "/".join(segments)


class TestSheetNames:
    """The bug this class exists for: a title longer than Excel's 31-character
    sheet-name limit is truncated by the workbook, and the connection builder
    -- which finds sheets BY NAME -- then found nothing. Every step reported
    success and the workbook came out with no connection in it at all."""

    def test_a_long_title_still_gets_its_connection(self):
        from app.export.sheets import assign_sheet_names

        title = "CUSTOMER_COUNT by MARKET_SEGMENT"  # 32 characters
        assert len(title) > 31
        name = assign_sheet_names([title])[0]
        assert name != title

        xlsx = build_workbook(
            "Sales",
            "D.S.V",
            "ALICE",
            [
                SheetData(
                    title=title,
                    columns=[{"name": "SEGMENT"}, {"name": "N"}],
                    rows=[["A", 1]],
                )
            ],
            generated_at=datetime(2026, 8, 17, tzinfo=timezone.utc),
        )
        out = add_live_connections(
            xlsx,
            [LiveSheet(title=name, sql=SQL, columns=["SEGMENT", "N"], rows=1)],
            account="A",
            database="D",
            schema="S",
        )
        assert "xl/connections.xml" in parts(out)

    def test_names_are_assigned_over_every_sheet_including_failed_ones(self):
        # A sheet that could not be read still occupies a name. Skipping it
        # when assigning would shift every later sheet's name by one, and the
        # connections would attach to the wrong tabs.
        from app.export.sheets import assign_sheet_names

        assert assign_sheet_names(["A", "B", "C"]) == ["A", "B", "C"]
        assert assign_sheet_names(["Same", "Same"]) == ["Same", "Same (2)"]
        assert assign_sheet_names(["Summary"]) == ["Summary (2)"]


class TestRelationshipNamespaces:
    """Excel called the workbook corrupt, and this is why.

    A .rels part has TWO namespaces in play. The `<Relationships>` container
    belongs to the PACKAGE namespace; the `Type` on each relationship inside
    belongs to the officeDocument one. Writing the container in the
    officeDocument namespace produces a file that is well-formed, whose every
    Target resolves, and which tells Excel the part declares no relationships
    at all -- so the worksheet's `r:id="rId1"` points at nothing and the
    package is rejected.

    The earlier tests all passed on that file, because they matched raw text
    with regexes. These parse.
    """

    PACKAGE = "http://schemas.openxmlformats.org/package/2006/relationships"

    def test_every_rels_container_is_in_the_package_namespace(self):
        from xml.etree import ElementTree

        archive = parts(live())
        rels = [n for n in archive if n.endswith(".rels")]
        assert rels
        for name in rels:
            root = ElementTree.fromstring(archive[name])
            assert root.tag == f"{{{self.PACKAGE}}}Relationships", name

    def test_every_r_id_a_worksheet_uses_resolves_in_its_own_rels(self):
        # Parsed rather than pattern-matched: a relationship in the wrong
        # namespace simply is not a relationship, and only a namespace-aware
        # reader notices.
        from xml.etree import ElementTree

        archive = parts(live())
        r_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        for name, blob in archive.items():
            if not (name.startswith("xl/worksheets/sheet") and name.endswith(".xml")):
                continue
            used = {
                el.get(r_ns)
                for el in ElementTree.fromstring(blob).iter()
                if el.get(r_ns)
            }
            if not used:
                continue
            rels_path = name.replace("xl/worksheets/", "xl/worksheets/_rels/") + ".rels"
            assert rels_path in archive, name
            declared = {
                el.get("Id")
                for el in ElementTree.fromstring(archive[rels_path]).findall(
                    f"{{{self.PACKAGE}}}Relationship"
                )
            }
            assert used <= declared, f"{name}: {used - declared} not declared"

    def test_the_table_finds_its_query_table(self):
        from xml.etree import ElementTree

        archive = parts(live())
        rels = ElementTree.fromstring(archive["xl/tables/_rels/table1.xml.rels"])
        targets = [
            r.get("Target")
            for r in rels.findall(f"{{{self.PACKAGE}}}Relationship")
        ]
        assert targets == ["../queryTables/queryTable1.xml"]

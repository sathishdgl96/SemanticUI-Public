import io

import openpyxl
import pytest
from snowflake.connector.errors import ProgrammingError

from app.auth.sessions import SESSION_COOKIE, create_session
from app.db.models import Report, Workspace, WorkspaceMember
from app.snowflake.provider import get_cache
from tests.test_report_routes import sign_in, valid_definition
from tests.test_semantic_routes import ScriptedConnection

SHEET = {
    "title": "Revenue by region",
    "dimensions": ["ORDERS.ORDER_DATE"],
    "metrics": ["ORDERS.TOTAL_REVENUE"],
}


def _report_for(db, sess, workspace):
    row = Report(
        owner_user_id=sess.user_id,
        workspace_id=workspace.id,
        name="Sales overview",
        view_database="ANALYTICS",
        view_schema="PUBLIC",
        view_name="SALES",
        definition=valid_definition(),
    )
    db.add(row)
    db.commit()
    return row


@pytest.fixture
def report(client, db):
    sess = sign_in(client, db)
    db.commit()
    workspace = (
        db.query(Workspace)
        .join(WorkspaceMember, WorkspaceMember.workspace_id == Workspace.id)
        .filter(WorkspaceMember.user_id == sess.user_id)
        .one()
    )
    get_cache().put(sess.id, ScriptedConnection())
    return _report_for(db, sess, workspace)


def export(client, report, sheets=None):
    return client.post(
        f"/api/reports/{report.id}/export.xlsx", json={"sheets": sheets or [SHEET]}
    )


def open_workbook(response):
    return openpyxl.load_workbook(io.BytesIO(response.content))


def test_export_requires_auth(client):
    assert (
        client.post("/api/reports/x/export.xlsx", json={"sheets": []}).status_code == 401
    )


def test_the_response_is_a_readable_workbook(client, db, report):
    response = export(client, report)
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "Sales overview.xlsx" in response.headers["content-disposition"]
    assert open_workbook(response).sheetnames == ["Summary", "Revenue by region"]


def test_the_sheet_carries_the_queried_rows(client, db, report):
    ws = open_workbook(export(client, report))["Revenue by region"]
    assert ws.cell(row=1, column=1).value == "ORDER_DATE"
    assert ws.cell(row=2, column=1).value == "2026-01-01"


def test_a_filter_is_bound_rather_than_interpolated(client, db, report):
    """Export runs the same validated path a normal query does."""
    conn = get_cache()
    response = export(
        client,
        report,
        sheets=[
            {
                **SHEET,
                "filters": [
                    {
                        "id": "f",
                        "field": "CUSTOMERS.REGION",
                        "op": "is",
                        "values": ["EAST"],
                    }
                ],
            }
        ],
    )
    assert response.status_code == 200
    entry = conn._entries[next(iter(conn._entries))]  # noqa: SLF001
    executed = entry.conn.cursor_obj.executed
    bound = entry.conn.cursor_obj.bound
    at = next(
        i
        for i, s in enumerate(executed)
        if "SEMANTIC_VIEW" in s and not s.startswith("DESCRIBE")
    )
    assert bound[at] == ["EAST"]
    assert "EAST" not in executed[at]


def test_the_context_reaches_the_summary_sheet(client, db, report):
    """Drill position travels with the export so the file explains itself."""
    response = export(
        client, report, sheets=[{**SHEET, "context": "Drilled into US > California"}]
    )
    workbook = open_workbook(response)
    text = "\n".join(
        str(c.value)
        for row in workbook["Summary"].iter_rows()
        for c in row
        if c.value is not None
    )
    assert "Drilled into US > California" in text


def test_a_forbidden_visual_does_not_lose_the_whole_workbook(client, db):
    """The colleague-with-narrower-permissions case."""

    class PartlyForbidden(ScriptedConnection):
        def cursor(self):
            cursor = super().cursor()
            if getattr(cursor, "_wrapped", False):
                return cursor
            original = cursor.execute

            def execute(sql, params=None):
                if "TOTAL_REVENUE" in sql and not sql.startswith("DESCRIBE"):
                    raise ProgrammingError(msg="Insufficient privileges", errno=3001)
                return original(sql, params)

            cursor.execute = execute
            cursor._wrapped = True
            return cursor

    sess = sign_in(client, db)
    db.commit()
    workspace = (
        db.query(Workspace)
        .join(WorkspaceMember, WorkspaceMember.workspace_id == Workspace.id)
        .filter(WorkspaceMember.user_id == sess.user_id)
        .one()
    )
    get_cache().put(sess.id, PartlyForbidden())
    row = _report_for(db, sess, workspace)

    response = client.post(
        f"/api/reports/{row.id}/export.xlsx",
        json={
            "sheets": [
                {"title": "OK", "dimensions": ["ORDERS.ORDER_DATE"], "metrics": []},
                {
                    "title": "Denied",
                    "dimensions": [],
                    "metrics": ["ORDERS.TOTAL_REVENUE"],
                },
            ]
        },
    )
    assert response.status_code == 200
    workbook = open_workbook(response)
    denied = "\n".join(
        str(c.value)
        for r in workbook["Denied"].iter_rows()
        for c in r
        if c.value is not None
    )
    assert "privileges" in denied.lower()
    assert workbook["OK"].cell(row=2, column=1).value == "2026-01-01"


def test_an_unknown_field_fails_the_whole_export(client, db, report):
    """A client bug, not a permissions difference -- so it is loud rather than
    degraded into a note on one sheet."""
    response = export(
        client, report, sheets=[{"title": "X", "dimensions": ["A.NOPE"], "metrics": []}]
    )
    assert response.status_code == 400


def test_exporting_a_report_i_cannot_see_is_404(client, db, report):
    other = create_session(db, account="ACME", user="BOB", mode="dev")
    db.commit()
    get_cache().put(other.id, ScriptedConnection())
    client.cookies.set(SESSION_COOKIE, other.id)
    assert export(client, report).status_code == 404


def test_too_many_sheets_is_rejected(client, db, report):
    assert export(client, report, sheets=[SHEET] * 51).status_code == 422


def test_a_report_name_that_would_break_the_header_is_sanitised(client, db):
    sess = sign_in(client, db)
    db.commit()
    workspace = (
        db.query(Workspace)
        .join(WorkspaceMember, WorkspaceMember.workspace_id == Workspace.id)
        .filter(WorkspaceMember.user_id == sess.user_id)
        .one()
    )
    get_cache().put(sess.id, ScriptedConnection())
    row = _report_for(db, sess, workspace)
    row.name = 'Q1/Q2 "report"\r\nX'
    db.commit()

    response = export(client, row)
    disposition = response.headers["content-disposition"]
    assert "\r" not in disposition and "\n" not in disposition
    assert '"' not in disposition.split("filename=")[1].strip('"')


class TestConnect:
    def test_it_returns_literal_sql_and_the_account(self, client, db, report):
        response = client.post(
            f"/api/reports/{report.id}/connect", json={"sheets": [SHEET]}
        )
        assert response.status_code == 200
        body = response.json()
        assert body["database"] == "ANALYTICS"
        assert body["view"] == "SALES"
        assert body["account"] == "ACME"
        assert "?" not in body["sheets"][0]["sql"]
        assert "SEMANTIC_VIEW" in body["sheets"][0]["sql"]

    def test_a_filter_value_is_inlined_for_power_query(self, client, db, report):
        response = client.post(
            f"/api/reports/{report.id}/connect",
            json={
                "sheets": [
                    {
                        **SHEET,
                        "filters": [
                            {
                                "id": "f",
                                "field": "CUSTOMERS.REGION",
                                "op": "is",
                                "values": ["EAST"],
                            }
                        ],
                    }
                ]
            },
        )
        sql = response.json()["sheets"][0]["sql"]
        assert "'EAST'" in sql
        assert "?" not in sql

    def test_it_needs_only_viewer(self, client, db):
        sess = sign_in(client, db)
        ws = Workspace(name="Team", kind="shared", snowflake_account="ACME")
        db.add(ws)
        db.flush()
        db.add(WorkspaceMember(workspace_id=ws.id, user_id=sess.user_id, role="viewer"))
        db.commit()
        get_cache().put(sess.id, ScriptedConnection())
        row = _report_for(db, sess, ws)
        response = client.post(
            f"/api/reports/{row.id}/connect", json={"sheets": [SHEET]}
        )
        assert response.status_code == 200

    def test_a_report_i_cannot_see_is_404(self, client, db, report):
        other = create_session(db, account="ACME", user="BOB", mode="dev")
        db.commit()
        get_cache().put(other.id, ScriptedConnection())
        client.cookies.set(SESSION_COOKIE, other.id)
        response = client.post(
            f"/api/reports/{report.id}/connect", json={"sheets": [SHEET]}
        )
        assert response.status_code == 404

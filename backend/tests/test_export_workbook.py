import io
from datetime import datetime, timezone

import openpyxl

from app.export.workbook import SheetData, build_workbook

WHEN = datetime(2026, 8, 16, 12, 0, tzinfo=timezone.utc)


def build(sheets, name="Sales overview"):
    """Build a workbook and read it back with a DIFFERENT library.

    openpyxl rather than trusting xlsxwriter's own view of what it wrote: the
    question is what Excel will see, and a reader sharing no code with the
    writer is the closest available proxy.
    """
    data = build_workbook(
        name, "ANALYTICS.PUBLIC.SALES", "ALICE", sheets, generated_at=WHEN
    )
    return openpyxl.load_workbook(io.BytesIO(data))


def a_sheet(**over):
    base = {
        "title": "Revenue by region",
        "columns": [{"name": "REGION"}, {"name": "TOTAL_REVENUE"}],
        "rows": [["EAST", 100], ["WEST", 200]],
    }
    return SheetData(**{**base, **over})


def summary_text(workbook):
    return "\n".join(
        str(c.value)
        for row in workbook["Summary"].iter_rows()
        for c in row
        if c.value is not None
    )


class TestFormulaInjection:
    """The thing this feature would otherwise get wrong."""

    PAYLOADS = [
        "=cmd|'/c calc'!A0",
        '=HYPERLINK("http://evil","click")',
        "+HYPERLINK(1)",
        "-2+3+cmd|' /C calc'!A0",
        "@SUM(1+1)",
    ]

    def test_a_dangerous_value_is_stored_as_text_not_a_formula(self):
        wb = build([a_sheet(rows=[[p, 1] for p in self.PAYLOADS])])
        ws = wb["Revenue by region"]
        for row, payload in enumerate(self.PAYLOADS, start=2):
            cell = ws.cell(row=row, column=1)
            assert cell.data_type == "s", f"{payload!r} was not stored as a string"
            assert cell.value == payload, "the displayed text must be unchanged"

    def test_the_value_is_not_prefixed_or_mangled(self):
        """Preferred over the usual apostrophe trick, which changes what the
        reader sees."""
        wb = build([a_sheet(rows=[["=1+1", 1]])])
        assert wb["Revenue by region"].cell(row=2, column=1).value == "=1+1"

    def test_a_negative_number_stays_a_number(self):
        """-5 begins with a dangerous prefix but is not text. Writing it as a
        string would break every sum in the sheet."""
        wb = build([a_sheet(rows=[["EAST", -5]])])
        cell = wb["Revenue by region"].cell(row=2, column=2)
        assert cell.value == -5
        assert cell.data_type == "n"

    def test_a_dangerous_value_in_a_header_is_also_text(self):
        """Column names come from the warehouse too."""
        wb = build([a_sheet(columns=[{"name": "=1+1"}, {"name": "OK"}])])
        cell = wb["Revenue by region"].cell(row=1, column=1)
        assert cell.data_type == "s"
        assert cell.value == "=1+1"


class TestShape:
    def test_a_summary_sheet_comes_first(self):
        assert build([a_sheet()]).sheetnames[0] == "Summary"

    def test_the_summary_names_the_report_the_view_and_who_exported_it(self):
        text = summary_text(build([a_sheet()]))
        assert "Sales overview" in text
        assert "ANALYTICS.PUBLIC.SALES" in text
        assert "ALICE" in text
        assert "2026-08-16" in text

    def test_the_summary_records_each_sheets_context(self):
        """A file found in a shared drive six months later should explain
        itself."""
        text = summary_text(build([a_sheet(context="Drilled into US > California")]))
        assert "Drilled into US > California" in text

    def test_the_summary_records_the_filters_applied(self):
        text = summary_text(build([a_sheet(filters=["CUSTOMERS.REGION is"])]))
        assert "CUSTOMERS.REGION" in text

    def test_one_sheet_per_visual_named_after_it(self):
        wb = build([a_sheet(title="A"), a_sheet(title="B")])
        assert wb.sheetnames == ["Summary", "A", "B"]

    def test_duplicate_titles_do_not_overwrite_each_other(self):
        wb = build([a_sheet(title="Revenue"), a_sheet(title="Revenue")])
        assert wb.sheetnames == ["Summary", "Revenue", "Revenue (2)"]

    def test_a_visual_titled_summary_does_not_collide_with_the_summary_sheet(self):
        wb = build([a_sheet(title="Summary")])
        assert wb.sheetnames == ["Summary", "Summary (2)"]

    def test_the_header_row_carries_the_column_names(self):
        ws = build([a_sheet()])["Revenue by region"]
        assert [ws.cell(row=1, column=i).value for i in (1, 2)] == [
            "REGION",
            "TOTAL_REVENUE",
        ]

    def test_the_rows_are_written_in_order(self):
        ws = build([a_sheet()])["Revenue by region"]
        assert ws.cell(row=2, column=1).value == "EAST"
        assert ws.cell(row=3, column=1).value == "WEST"

    def test_a_null_becomes_an_empty_cell_not_the_text_None(self):
        ws = build([a_sheet(rows=[[None, 1]])])["Revenue by region"]
        assert ws.cell(row=2, column=1).value in (None, "")


class TestFailureAndLimits:
    def test_a_failed_sheet_carries_its_error_and_the_others_still_carry_data(self):
        """A colleague with narrower Snowflake permissions gets a partial
        workbook that explains itself, not a 500."""
        wb = build(
            [
                a_sheet(title="Allowed"),
                a_sheet(title="Denied", rows=[], error="Insufficient privileges"),
            ]
        )
        assert wb["Allowed"].cell(row=2, column=1).value == "EAST"
        denied = "\n".join(
            str(c.value)
            for row in wb["Denied"].iter_rows()
            for c in row
            if c.value is not None
        )
        assert "Insufficient privileges" in denied

    def test_the_summary_names_which_sheets_failed(self):
        text = summary_text(build([a_sheet(title="Denied", rows=[], error="nope")]))
        assert "Denied" in text
        assert "nope" in text

    def test_truncation_is_declared_rather_than_silent(self, monkeypatch):
        from app.export import workbook as workbook_module

        monkeypatch.setattr(workbook_module, "MAX_ROWS_PER_SHEET", 2)
        wb = build([a_sheet(rows=[["A", 1], ["B", 2], ["C", 3]])])
        assert wb["Revenue by region"].max_row == 3, "header plus two rows"
        assert "truncat" in summary_text(wb).lower()

    def test_an_empty_result_is_a_sheet_with_only_headers(self):
        ws = build([a_sheet(rows=[])])["Revenue by region"]
        assert ws.cell(row=1, column=1).value == "REGION"
        assert ws.max_row == 1

    def test_a_workbook_with_no_sheets_at_all_still_opens(self):
        wb = build([])
        assert wb.sheetnames == ["Summary"]


def test_the_output_is_a_readable_workbook_not_merely_bytes():
    data = build_workbook("R", "V", "ALICE", [a_sheet()], generated_at=WHEN)
    assert data[:2] == b"PK", "an xlsx is a zip archive"
    openpyxl.load_workbook(io.BytesIO(data))
